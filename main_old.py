from docx import Document
from openpyxl import load_workbook
from calendar import monthrange, weekday
from teachers_list import teachers_list
from glob import glob

import time
import shutil

def resave_file(src: str, dst='tabel.xlsx') -> None:
    shutil.copy2(src=src, dst=dst)

def script(gen_excel_table_path: str, document_path: str) -> None:
    start_values = get_start_day(document_path=document_path)

    days_amount = start_values[0]
    week_num = start_values[1]
    day_num = start_values[-1]
    date = start_values[2].split('.')[1:]

    resave_file(src=f'resources/tables/base{days_amount}.xlsx')

    wb_gen = load_workbook(filename=gen_excel_table_path)
    ws_gen = wb_gen.active

    wb_base = load_workbook(filename='tabel.xlsx')
    ws_base = wb_base.active

    last_full_column = 2

    for teacher in teachers_list:
        ws_base._get_cell(row=1, column=last_full_column+1).value = teacher
        last_full_column += 1
    else:
        wb_base.save(filename='tabel.xlsx')

    curr_str = 2

    for i in range(days_amount):
        ws_base[f'A{curr_str}'].value = get_day(date=f'{i+1}.{date[0]}.{date[1]}', week_num=week_num)

        if 'воскресенье' in ws_base[f'A{curr_str}'].value:
            if week_num == '1':
                week_num = '2'
            else:
                week_num = '1'

        for teacher_num in range(len(teachers_list)):
            teacher = teachers_list[teacher_num].split()[0]
            if 'суббота' not in ws_base[f'A{curr_str}'].value and 'воскресенье' not in ws_base[f'A{curr_str}'].value:
                for row in range(days_strs[week_days[day_num]], days_strs[week_days[day_num+1]]):
                    #перебираем все номера столбцов в таблице
                    for col in range(1, ws_gen.max_column+1):
                        #получение ячейки: row-№строки, col-№столбца
                        cell = ws_gen._get_cell(row=row, column=col)
                        
                        #получение значения ячейки
                        value = str(cell.value)
                        value_ed = del_all_spaces(value)

                        #проверка наличия ключевого слова в ячейке (фамилия преподавателя)
                        if teacher in value.split():
                            #время пары
                            time = del_all_spaces(ws_gen._get_cell(row=cell.row, column=2).value)
                            #название группы
                            group = get_group_name(get_group(del_spaces(ws_gen._get_cell(row=cell.row-1-offset[time], column=cell.column).value)))
                            #день недели
                            #day = del_spaces(ws_gen._get_cell(row=cell.row-offset[time], column=1).value)
                            #определяет, есть ли разделение пар по неделям
                            curr_cell = ws_base._get_cell(row=curr_str+offset[time], column=teacher_num+3)
                            if '1н' in value_ed or '2н' in value_ed:
                                values = del_spaces(value).split('2 н')
                                if len(values) == 1:
                                    values = value.split('2н')
                                #если пара с номером недели в расписании не соответсвует номеру нужной недели, изменения не вносятся, ячейка пропускается
                                if (teacher not in values[0] and week_num == '1') or (teacher not in values[1] and week_num == '2'):
                                    continue
                            
                            if curr_cell.value == None:
                                curr_cell.value = f'{group}\n'
                            else:
                                curr_cell.value = curr_cell.value+f'{group}\n'
            if i+1 < 10:
                curr_date = f'0{i+1}.{date[0]}.{date[1]}'
            else:
                curr_date = f'{i+1}.{date[0]}.{date[1]}'
            if curr_date in doc_tables:
                for doc_row in doc_tables[curr_date].rows:
                    cells_list = doc_row.cells
                    text_3 = del_spaces(cells_list[3].text)
                    text_5 = del_spaces(cells_list[5].text)

                    if teacher in text_3.split() and text_3 != text_5:
                        #получаем группу из word документа для которой есть замена 
                        group_curr = f'{cells_list[1].text.split()[0]} - {cells_list[1].text.split()[-1]}'
                        #получаем список групп по расписанию из конечной таблицы
                        curr_cell_w = ws_base._get_cell(row=curr_str+int(cells_list[0].text)-1, column=teacher_num+3)
                        group_list = curr_cell_w.value

                        if group_list != None:
                            group_list = group_list.split('\n')
                        else:
                            group_list = ''

                        #убираем группу из расписания
                        if group_curr in group_list:
                            group_list.remove(group_curr)
                            text = '\n'.join(group for group in group_list if group != '')
                            ws_base._get_cell(row=curr_str+int(cells_list[0].text)-1, column=teacher_num+3).value = f'{text}\n'

                    elif teacher in text_5.split() and text_3 != text_5:
                        #получаем группу из word документа для которой есть замена 
                        group_curr = f'{cells_list[1].text.split()[0]} - {cells_list[1].text.split()[-1]}'
                        #получаем список групп по расписанию из конечной таблицы
                        if '-' not in cells_list[0].text:
                            curr_cell_w = ws_base._get_cell(row=curr_str+int(cells_list[0].text)-1, column=teacher_num+3)
                            group_list = curr_cell_w.value
                        else:
                            curr_cell_w = ws_base._get_cell(row=curr_str+int(cells_list[0].text.split('-')[0])-1, column=teacher_num+3)
                            group_list = curr_cell_w.value

                        if group_list != None and '-' not in cells_list[0].text:
                            ws_base._get_cell(row=curr_str+int(cells_list[0].text)-1, column=teacher_num+3).value = group_list+f'{group_curr}\n'
                        elif '-' in cells_list[0].text:
                            pass
                        else:
                            ws_base._get_cell(row=curr_str+int(cells_list[0].text)-1, column=teacher_num+3).value = f'{group_curr}\n'

        if day_num == 6:
            day_num = 0
        else:
            day_num += 1

        curr_str += 6
    
    else:
        wb_base.save(filename='tabel.xlsx')
        wb_base.close()
        wb_gen.close()


def get_start_day(document_path: str):
    doc = Document(document_path)

    for num in paragraphs:
        text = doc.paragraphs[num].text.split()
        date = text[0]

        if '01' in date:
            days = monthrange(year=int(date.split('.')[-1]), month=int(date.split('.')[1]))
            return days[1], text[1], date, days[0]

def get_day(date: str, week_num: str):
    day = week_days[weekday(year=int(date.split('.')[-1]), month=int(date.split('.')[1]), day=int(date.split('.')[0]))]
    return f'{date}\n({day} {week_num}н)'

#удаление лишних пробелов из строки
def del_spaces(string: str) -> str:
    if string != None:
        return ' '.join([str(i) for i in string.split()])

#удаление всех пробелов из строки
def del_all_spaces(string: str) -> str:
    if string != None:
        return string.replace(' ', '')

#удаление лишних символов из названия группы (из общей таблицы с расписанием)
def get_group(group_name: str) -> str:
    group = group_name.split('(')[-1]
    return group[:-1]

def get_group_name(group_name:str) -> str:
    if '-' not in group_name:
        return f'{group_name.split()[0]} - {group_name.split()[1]}'
    else:
        return group_name


def get_doc_tables_dict() -> dict:
    files_amount = len(glob(r'resources\docs\*.docx'))

    for doc_num in range(files_amount):
        doc = Document(f'resources\docs\{doc_num+1}.docx')
        tables = doc.tables

        if len(tables) >= 6:
            tables = tables[:-1]

        for table_num in range(len(tables)):
            table = tables[table_num]
            date = doc.paragraphs[paragraphs[table_num]].text.split()[0]
            doc_tables[date] = table
        

def main():
    s = time.time()
    get_doc_tables_dict()
    #print(doc_tables)
    script(gen_excel_table_path=r'resources\tables\RASPISANIE.xlsx', document_path=r'resources\docs\1.docx')
    print(time.time()-s)


offset = {'08.45-10.15': 0, '10.30-12.00': 1, '12.40-14.10': 2, '14.20-15.50': 3, '16.00-17.30': 4, '17.40-19.10': 5}
paragraphs = [3, 10, 18, 26, 34]
week_days = {0: 'понедельник', 1: 'вторник', 2: 'среда', 3: 'четверг', 4: 'пятница', 5: 'суббота', 6: 'воскресенье'}
days_strs = {'понедельник': 13, 'вторник': 20, 'среда': 26, 'четверг': 32, 'пятница': 39, 'суббота': 45}
doc_tables = {}


main()